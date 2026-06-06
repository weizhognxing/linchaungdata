import json
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from flask import session
from openpyxl import Workbook

from app.common import required
from app.db import db
from app.services.core import ask_ai_yes_no, get_fields_for_disease


# Shared constants and helper functions for app.routes.records.
# Keep route handlers in records.py; put reusable query/normalization logic here.

__all__ = [
    "DIAGNOSIS_DISEASE_OPTIONS",
    "MEDICAL_HISTORY_OPTIONS",
    "DIAGNOSIS_SUBCATEGORY_OPTIONS",
    "LAB_CATEGORY_DEFINITIONS",
    "INTAKE_PROMPT",
    "_current_user_scope",
    "_patient_columns",
    "_table_exists",
    "_xlsx_value",
    "_write_sheet",
    "_build_single_sheet_workbook",
    "_fetch_rows_for_patients",
    "_normalize_medical_history",
    "_normalize_preliminary_diagnosis",
    "_optional_number",
    "_optional_decimal",
    "_optional_datetime",
    "_request_flag_enabled",
    "_lab_category_config",
    "_build_category_prompt",
    "_fields_for_category",
    "_patient_is_accessible",
    "_patient_is_complete",
    "_missing_case_completion_items",
    "_required_lab_tests_completed",
    "_mark_case_complete_if_ready",
    "_find_existing_patient_for_lab",
    "_find_duplicate_lab_record",
]

DIAGNOSIS_DISEASE_OPTIONS = {"脓毒症", "重症肺炎", "ARDS", "心肺复苏后", "急性坏死性胰腺炎", "消化道出血", "中毒", "心源性休克/心衰", "脑卒中", "多发伤", "颅脑损伤", "胸部损伤", "热射病"}
MEDICAL_HISTORY_OPTIONS = {"无", "糖尿病", "乙肝/肝硬化", "肾衰", "肝衰", "冠心病", "COPD", "高血压"}
DIAGNOSIS_SUBCATEGORY_OPTIONS = {
    "脓毒症": {"肺部", "腹部", "心血管/血液", "泌尿系", "脑部", "软组织", "不详"},
    "重症肺炎": set(),
    "ARDS": set(),
    "心肺复苏后": set(),
    "急性坏死性胰腺炎": set(),
    "消化道出血": set(),
    "中毒": {"有机磷中毒", "CO中毒", "蘑菇中毒", "杀虫剂/除草剂中毒", "药物中毒"},
    "心源性休克/心衰": {"1心肌梗塞", "2心衰", "3.心肌炎", "4.急性瓣膜病变", "5电传导病变"},
    "脑卒中": {"脑梗塞", "基底节出血", "小脑出血", "蛛网膜下腔出血", "脑干出血"},
    "多发伤": {"颅脑损伤", "胸部创伤", "腹部创伤", "四肢损伤", "脊柱损伤"},
    "颅脑损伤": {"大脑挫裂伤", "缺氧缺血性脑病", "弥漫性轴索损伤"},
    "胸部损伤": {"连枷胸", "开放性气胸", "三根以上肋骨骨折", "开放性血气胸"},
    "热射病": set(),
}

LAB_CATEGORY_DEFINITIONS = {
    "blood_routine": {
        "label": "血细胞分析",
        "fields": [
            "wbc", "neu", "lym", "mono", "eos", "baso", "neu_r", "lym_r", "mono_r", "eos_r", "baso_r",
            "rbc", "hgb", "hct", "mcv", "mch", "mchc", "rdw_sd", "rdw_cv", "plt", "mpv", "pdw", "p_lcr", "hscrp",
        ],
    },
    "biochemistry": {
        "label": "生化电解质",
        "fields": [
            "alt", "ast", "ast_alt", "tp", "alb", "glo", "a_g", "tbil", "dbil", "ibil", "tba", "ggt", "alp", "pa",
            "urea", "ua", "crea", "k", "na", "cl", "ca", "co2", "p", "mg", "fe", "gfr", "ag", "hemo", "icte", "lipe",
        ],
    },
    "blood_gas": {
        "label": "血气分析",
        "fields": ["ph", "pco2", "po2", "k", "na", "cl", "ca__", "glu", "lac_la", "thb", "so2", "o2hb"],
    },
    "dic7": {
        "label": "DIC7项",
        "fields": ["pt", "pt_inr", "pta", "pt_ratio", "aptt", "aptt_r", "tt", "fib", "d_di", "fdp", "at3"],
    },
    "bnp": {
        "label": "BNP",
        "fields": ["nt_probnp", "hs_ctnt_stat", "myo_stat", "ck_mb_stat"],
    },
    "lactate": {
        "label": "乳酸",
        "fields": ["lactate", "lac_la"],
    },
    "pct": {
        "label": "PCT",
        "fields": ["pct"],
    },
}

INTAKE_PROMPT = """请识别这张检验图片或病历图片中的患者基础信息，仅返回基础信息，不要提取检验指标。
请严格按照以下 JSON 格式返回，只返回 JSON，不要其他内容：
{
  "lab_test_name": "患者基础信息",
  "patient": {
    "name": "姓名",
    "gender": "性别",
    "age": "年龄",
    "phone": "",
    "id_number": ""
  },
  "items": []
}"""


def _current_user_scope(cur):
    user_id = session.get("user_id")
    cur.execute("SELECT id, hospital_id FROM users WHERE id=%s", (user_id,))
    return cur.fetchone()


def _patient_columns(cur):
    cur.execute("SHOW COLUMNS FROM patients")
    return {row["Field"] for row in cur.fetchall()}


def _table_exists(cur, table_name):
    cur.execute("SHOW TABLES LIKE %s", (table_name,))
    return cur.fetchone() is not None


def _xlsx_value(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d %H:%M:%S") if isinstance(value, datetime) else value.strftime("%Y-%m-%d")
    return value


def _write_sheet(workbook, title, rows, headers):
    sheet = workbook.create_sheet(title=title)
    sheet.append(headers)
    for row in rows:
        sheet.append([_xlsx_value(row.get(header)) for header in headers])
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 10), 36)


def _build_single_sheet_workbook(sheet_title, rows, headers):
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(workbook, sheet_title, rows, headers)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()


def _fetch_rows_for_patients(cur, table_name, patient_ids, patient_column="patient_id"):
    if not _table_exists(cur, table_name):
        return [], []
    cur.execute(f"SHOW COLUMNS FROM `{table_name}`")
    headers = [row["Field"] for row in cur.fetchall()]
    if not patient_ids:
        return [], headers
    placeholders = ",".join(["%s"] * len(patient_ids))
    cur.execute(
        f"SELECT * FROM `{table_name}` WHERE `{patient_column}` IN ({placeholders}) ORDER BY `{patient_column}`, id",
        patient_ids,
    )
    return cur.fetchall(), headers


def _normalize_medical_history(raw_value):
    if isinstance(raw_value, list):
        items = raw_value
    else:
        items = str(raw_value or "无").split(",")
    selected = []
    for item in items:
        value = str(item).strip()
        if value in MEDICAL_HISTORY_OPTIONS and value not in selected:
            selected.append(value)
    if not selected or "无" in selected:
        return "无"
    return ",".join(selected)


def _normalize_preliminary_diagnosis(disease, raw_value):
    if not disease:
        return ""
    allowed = DIAGNOSIS_SUBCATEGORY_OPTIONS.get(disease, set())
    selected = []
    for item in str(raw_value or "").split(","):
        value = item.strip()
        if value in allowed and value not in selected:
            selected.append(value)
    if disease == "多发伤":
        selected = selected[:3]
    elif len(selected) > 1:
        selected = selected[:1]
    return ",".join(selected)


def _optional_number(value):
    value = str(value or "").strip()
    return value if value != "" else None


def _optional_decimal(value):
    value = str(value or "").strip()
    return value if value != "" else None


def _optional_datetime(value):
    value = str(value or "").strip()
    if value and "T" in value:
        value = value.replace("T", " ")
    if value and len(value) == 16:
        value += ":00"
    return value if value != "" else None


def _request_flag_enabled(data, key, default=True):
    if key not in data:
        return default
    value = data.get(key)
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() not in {"0", "false", "no", "off", ""}


def _lab_category_config(category):
    return LAB_CATEGORY_DEFINITIONS.get(str(category or "").strip())


def _build_category_prompt(category_config):
    fields = ", ".join(category_config["fields"])
    return f"""请识别这张检验报告图片，提取患者信息、图片上的实际检验名称和所有检验指标结果。
用户可能正在补录“{category_config['label']}”，但图片不一定属于这个类别，请不要强行归类。
如果图片确实属于“{category_config['label']}”，请优先使用以下字段缩写作为 code：{fields}
如果图片属于其他检验，请按图片原文或最接近的报告标题填写 lab_test_name，并提取图片中的实际检验指标。
lab_test_name 只能填写检验项目/报告名称本身，必须去掉前面的套餐编号、医嘱编号、条码编号、内部代码和序号，例如 JY243.血淀粉酶+P-AMY 应返回 血淀粉酶+P-AMY，A001-血细胞分析 应返回 血细胞分析。
请严格按照以下JSON格式返回，只返回JSON，不要其他内容：
{{
  "lab_test_name": "图片上的实际检验名称",
  "patient": {{
    "name": "姓名",
    "gender": "性别",
    "age": "年龄",
    "phone": "电话",
    "id_number": "登记号"
  }},
  "items": [
    {{"name": "检验指标中文名称", "code": "英文缩写", "value": "结果值", "unit": "单位", "reference": "参考区间", "abnormal": "↑/↓/正常"}}
  ]
}}"""


def _fields_for_category(fields, category):
    category_config = _lab_category_config(category)
    if not category_config:
        return fields
    field_names = set(category_config["fields"])
    filtered = [field for field in fields if field.get("field_name") in field_names]
    return filtered or fields


def _patient_is_accessible(cur, patient_id, current_user, patient_columns=None):
    patient_columns = patient_columns or _patient_columns(cur)
    if "hospital_id" in patient_columns:
        cur.execute(
            """
            SELECT p.id
            FROM patients p
            LEFT JOIN lab_records r ON r.patient_id=p.id
            LEFT JOIN users u ON u.id=r.user_id
            WHERE p.id=%s AND (p.hospital_id=%s OR u.hospital_id=%s)
            LIMIT 1
            """,
            (patient_id, current_user["hospital_id"], current_user["hospital_id"]),
        )
    else:
        cur.execute(
            """
            SELECT p.id
            FROM patients p
            JOIN lab_records r ON r.patient_id=p.id
            JOIN users u ON u.id=r.user_id
            WHERE p.id=%s AND u.hospital_id=%s
            LIMIT 1
            """,
            (patient_id, current_user["hospital_id"]),
        )
    return cur.fetchone() is not None


def _patient_is_complete(cur, patient_id, patient_columns=None):
    patient_columns = patient_columns or _patient_columns(cur)
    status_fields = [field for field in ("case_status", "case_integrity") if field in patient_columns]
    if not status_fields:
        return False
    cur.execute(f"SELECT {','.join(status_fields)} FROM patients WHERE id=%s LIMIT 1", (patient_id,))
    patient = cur.fetchone() or {}
    return patient.get("case_status") == "submitted" or patient.get("case_integrity") in {"complete", "submitted"}


def _missing_case_completion_items(cur, patient_id):
    cur.execute("SELECT name, gender, age FROM patients WHERE id=%s LIMIT 1", (patient_id,))
    patient = cur.fetchone() or {}
    missing = []
    if required(patient, ["name", "gender", "age"]):
        missing.append("基础信息")

    required_tables = [
        ("diagnosis_records", "诊断"),
        ("assessments", "评估"),
        ("followups", "随访"),
        ("treatments", "治疗"),
    ]
    for table_name, label in required_tables:
        if not _table_exists(cur, table_name):
            missing.append(label)
            continue
        cur.execute(f"SELECT COUNT(*) AS total FROM `{table_name}` WHERE patient_id=%s", (patient_id,))
        count_row = cur.fetchone() or {}
        if int(count_row.get("total") or 0) < 1:
            missing.append(label)
    if not _required_lab_tests_completed(cur, patient_id):
        missing.append("检验")
    return missing


def _required_lab_tests_completed(cur, patient_id):
    if not _table_exists(cur, "lab_records"):
        return False
    cur.execute("SHOW COLUMNS FROM lab_records")
    columns = {row["Field"] for row in cur.fetchall()}
    select_parts = []
    if "lab_test_name" in columns:
        select_parts.append("lab_test_name")
    if "record_category" in columns:
        select_parts.append("record_category")
    if not select_parts:
        return False

    cur.execute(
        f"SELECT {','.join(select_parts)} FROM lab_records WHERE patient_id=%s ORDER BY created_at, id",
        (patient_id,),
    )
    names = []
    for row in cur.fetchall():
        lab_test_name = str(row.get("lab_test_name") or "").strip()
        record_category = str(row.get("record_category") or "").strip()
        category_label = (LAB_CATEGORY_DEFINITIONS.get(record_category) or {}).get("label", "")
        if lab_test_name:
            names.append(lab_test_name)
        if category_label and category_label != lab_test_name:
            names.append(category_label)
    if not names:
        return False

    prompt = """你是临床检验类别判断助手。请判断下面“已完成的检验类别”中，是否已经覆盖全部七类必需检验类别：血细胞分析、生化电解质、血气分析、DIC7项、BNP、乳酸、PCT。
判断规则：
1. 名称可能带有编码或前缀，例如“JY2625.B型钠尿肽”，应理解为 BNP。
2. BNP 的同义名称包括：BNP、B型钠尿肽、N端B型钠尿肽、NT-proBNP、NT-pro BNP。
3. PCT 的同义名称包括：PCT、降钙素原。
4. 乳酸的同义名称包括：乳酸、LAC、LA。
5. DIC7项可写作：DIC7项、DIC 7项、凝血、凝血功能、凝血七项。
6. 只有七类都已覆盖，才返回“是”；只要缺少任意一类，就返回“否”。
7. 只能返回一个字：是 或 否，不要解释。

已完成的检验名称：
""" + "\n".join(f"- {name}" for name in names)
    try:
        return ask_ai_yes_no(prompt, timeout=60)
    except Exception as e:
        print(f"Lab completion AI check error: {e}")
        return False


def _mark_case_complete_if_ready(cur, patient_id, patient_columns=None):
    patient_columns = patient_columns or _patient_columns(cur)
    if "case_status" not in patient_columns and "case_integrity" not in patient_columns:
        return False

    status_fields = []
    if "case_status" in patient_columns:
        status_fields.append("case_status")
    if "case_integrity" in patient_columns:
        status_fields.append("case_integrity")
    cur.execute(
        f"SELECT {','.join(status_fields)} FROM patients WHERE id=%s LIMIT 1",
        (patient_id,),
    )
    patient = cur.fetchone() or {}
    if patient.get("case_status") == "submitted" or patient.get("case_integrity") == "complete":
        return False
    if _missing_case_completion_items(cur, patient_id):
        return False

    updates = []
    if "case_status" in patient_columns:
        updates.append("case_status='submitted'")
    if "case_integrity" in patient_columns:
        updates.append("case_integrity='complete'")
    if "updated_at" in patient_columns:
        updates.append("updated_at=NOW()")
    cur.execute(f"UPDATE patients SET {','.join(updates)} WHERE id=%s", (patient_id,))
    return True


def _find_existing_patient_for_lab(cur, current_user, patient, patient_columns):
    conditions = ["name=%s", "IFNULL(gender,'')=%s", "IFNULL(age,'')=%s"]
    params = [patient.get("name"), patient.get("gender", ""), patient.get("age") or ""]
    if "hospital_id" in patient_columns:
        conditions.insert(0, "hospital_id=%s")
        params.insert(0, current_user["hospital_id"])
    cur.execute(
        f"SELECT id FROM patients WHERE {' AND '.join(conditions)} ORDER BY id LIMIT 1",
        params,
    )
    row = cur.fetchone()
    return row["id"] if row else None


def _find_duplicate_lab_record(cur, patient_id, existing_columns, record_category, lab_test_name, photo_path, record_values):
    if photo_path and "photo_path" in existing_columns:
        cur.execute(
            "SELECT id FROM lab_records WHERE patient_id=%s AND photo_path=%s LIMIT 1",
            (patient_id, photo_path),
        )
        row = cur.fetchone()
        if row:
            return row["id"]

    conditions = ["patient_id=%s"]
    params = [patient_id]
    if "record_category" in existing_columns:
        conditions.append("IFNULL(record_category,'')=%s")
        params.append(record_category or "")
    if "lab_test_name" in existing_columns:
        conditions.append("IFNULL(lab_test_name,'')=%s")
        params.append(lab_test_name or "")
    for field, value in record_values.items():
        conditions.append(f"IFNULL(`{field.replace('`', '``')}`,'')=%s")
        params.append(value or "")
    cur.execute(
        f"SELECT id FROM lab_records WHERE {' AND '.join(conditions)} ORDER BY id LIMIT 1",
        params,
    )
    row = cur.fetchone()
    return row["id"] if row else None


