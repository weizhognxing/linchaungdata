import json
import os
import time
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from flask import Blueprint, current_app, request, send_file, session
from openpyxl import Workbook
from werkzeug.utils import secure_filename

from app.common import fail, ok, required
from app.db import db
from app.decorators import require_user
from app.services.core import ask_ai_yes_no, get_fields_for_disease, parse_ai_result, parse_lab_test_name, parse_patient_info, recognize_image, reorder_fields_by_values


records_bp = Blueprint("records", __name__)

DIAGNOSIS_DISEASE_OPTIONS = {"脓毒症部位", "重症胰腺炎", "心源性休克/心脏骤停", "中毒", "脑损伤", "多发伤", "胸部创伤"}
MEDICAL_HISTORY_OPTIONS = {"无", "糖尿病", "乙肝/肝硬化", "肾衰", "肝衰", "冠心病", "COPD", "高血压"}
DIAGNOSIS_SUBCATEGORY_OPTIONS = {
    "脓毒症部位": {"肺部", "腹部", "心血管/血液", "泌尿系", "脑部", "软组织", "不详"},
    "重症胰腺炎": set(),
    "心源性休克/心脏骤停": {"1心肌梗塞", "2心衰", "3.心肌炎", "4.急性瓣膜病变", "5电传导病变"},
    "中毒": {"有机磷中毒", "CO中毒", "蘑菇中毒", "杀虫剂/除草剂中毒", "药物中毒"},
    "脑损伤": {"大脑挫裂伤", "缺氧缺血性脑病", "弥漫性轴索损伤", "基底节出血", "小脑出血", "蛛网膜下腔出血", "脑梗塞", "脑干出血", "热射病"},
    "多发伤": {"颅脑损伤", "胸部创伤", "腹部创伤", "四肢损伤", "脊柱损伤"},
    "胸部创伤": {"连枷胸", "开放性气胸", "三根以上肋骨骨折", "开放性血气胸"},
}

LAB_CATEGORY_DEFINITIONS = {
    "blood_routine": {
        "label": "血常规",
        "fields": [
            "wbc", "neu", "lym", "mono", "eos", "baso", "neu_r", "lym_r", "mono_r", "eos_r", "baso_r",
            "rbc", "hgb", "hct", "mcv", "mch", "mchc", "rdw_sd", "rdw_cv", "plt", "mpv", "pdw", "p_lcr", "hscrp",
        ],
    },
    "biochemistry": {
        "label": "生化电解质",
        "fields": [
            "alt", "ast", "ast_alt", "tp", "alb", "glo", "a_g", "tbil", "dbil", "ibil", "tba", "ggt", "alp", "pa",
            "urea", "ua", "crea", "k", "na", "cl", "ca", "co2", "p", "mg", "fe", "gfr", "ag", "hemo", "icte", "lipe",
        ],
    },
    "blood_gas": {
        "label": "血气分析",
        "fields": ["ph", "pco2", "po2", "k", "na", "cl", "ca__", "glu", "lac_la", "thb", "so2", "o2hb"],
    },
    "dic7": {
        "label": "DIC7项",
        "fields": ["pt", "pt_inr", "pta", "pt_ratio", "aptt", "aptt_r", "tt", "fib", "d_di", "fdp", "at3"],
    },
    "bnp": {
        "label": "BNP",
        "fields": ["nt_probnp", "hs_ctnt_stat", "myo_stat", "ck_mb_stat"],
    },
    "lactate": {
        "label": "乳酸",
        "fields": ["lactate", "lac_la"],
    },
    "pct": {
        "label": "PCT",
        "fields": ["pct"],
    },
}

INTAKE_PROMPT = """请识别这张检验图片或病历图片中的患者基础信息，仅返回基础信息，不要提取检验指标。
请严格按照以下 JSON 格式返回，只返回 JSON，不要其他内容：
{
  "lab_test_name": "本次图片对应的检验类别，例如患者基础信息/入院记录",
  "patient": {
    "name": "姓名",
    "gender": "性别",
    "age": "年龄",
    "phone": "电话",
    "id_number": "登记号"
  },
  "items": []
}"""


def _current_user_scope(cur):
    user_id = session.get("user_id")
    cur.execute("SELECT id, hospital_id FROM users WHERE id=%s", (user_id,))
    return cur.fetchone()


def _patient_columns(cur):
    cur.execute("SHOW COLUMNS FROM patients")
    return {row["Field"] for row in cur.fetchall()}


def _table_exists(cur, table_name):
    cur.execute("SHOW TABLES LIKE %s", (table_name,))
    return cur.fetchone() is not None


def _xlsx_value(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d %H:%M:%S") if isinstance(value, datetime) else value.strftime("%Y-%m-%d")
    return value


def _write_sheet(workbook, title, rows, headers):
    sheet = workbook.create_sheet(title=title)
    sheet.append(headers)
    for row in rows:
        sheet.append([_xlsx_value(row.get(header)) for header in headers])
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 10), 36)


def _fetch_rows_for_patients(cur, table_name, patient_ids, patient_column="patient_id"):
    if not _table_exists(cur, table_name):
        return [], []
    cur.execute(f"SHOW COLUMNS FROM `{table_name}`")
    headers = [row["Field"] for row in cur.fetchall()]
    if not patient_ids:
        return [], headers
    placeholders = ",".join(["%s"] * len(patient_ids))
    cur.execute(
        f"SELECT * FROM `{table_name}` WHERE `{patient_column}` IN ({placeholders}) ORDER BY `{patient_column}`, id",
        patient_ids,
    )
    return cur.fetchall(), headers


def _normalize_medical_history(raw_value):
    if isinstance(raw_value, list):
        items = raw_value
    else:
        items = str(raw_value or "无").split(",")
    selected = []
    for item in items:
        value = str(item).strip()
        if value in MEDICAL_HISTORY_OPTIONS and value not in selected:
            selected.append(value)
    if not selected or "无" in selected:
        return "无"
    return ",".join(selected)


def _normalize_preliminary_diagnosis(disease, raw_value):
    if not disease:
        return ""
    allowed = DIAGNOSIS_SUBCATEGORY_OPTIONS.get(disease, set())
    selected = []
    for item in str(raw_value or "").split(","):
        value = item.strip()
        if value in allowed and value not in selected:
            selected.append(value)
    if disease == "多发伤":
        selected = selected[:3]
    elif len(selected) > 1:
        selected = selected[:1]
    return ",".join(selected)


def _optional_number(value):
    value = str(value or "").strip()
    return value if value != "" else None


def _optional_decimal(value):
    value = str(value or "").strip()
    return value if value != "" else None


def _optional_datetime(value):
    value = str(value or "").strip()
    return value if value != "" else None


def _lab_category_config(category):
    return LAB_CATEGORY_DEFINITIONS.get(str(category or "").strip())


def _build_category_prompt(category_config):
    fields = ", ".join(category_config["fields"])
    return f"""请识别这张检验报告图片，提取患者信息、图片上的实际检验名称和所有检验指标结果。
用户可能正在补录“{category_config['label']}”，但图片不一定属于这个类别，请不要强行归类。
如果图片确实属于“{category_config['label']}”，请优先使用以下字段缩写作为 code：{fields}
如果图片属于其他检验，请按图片原文或最接近的报告标题填写 lab_test_name，并提取图片中的实际检验指标。
lab_test_name 只能填写检验项目/报告名称本身，必须去掉前面的套餐编号、医嘱编号、条码编号、内部代码和序号，例如 JY243.血淀粉酶+P-AMY 应返回 血淀粉酶+P-AMY，A001-血常规 应返回 血常规。
请严格按照以下JSON格式返回，只返回JSON，不要其他内容：
{{
  "lab_test_name": "图片上的实际检验名称",
  "patient": {{
    "name": "姓名",
    "gender": "性别",
    "age": "年龄",
    "phone": "电话",
    "id_number": "登记号"
  }},
  "items": [
    {{"name": "检验指标中文名称", "code": "英文缩写", "value": "结果值", "unit": "单位", "reference": "参考区间", "abnormal": "↑/↓/正常"}}
  ]
}}"""


def _fields_for_category(fields, category):
    category_config = _lab_category_config(category)
    if not category_config:
        return fields
    field_names = set(category_config["fields"])
    filtered = [field for field in fields if field.get("field_name") in field_names]
    return filtered or fields


def _patient_is_accessible(cur, patient_id, current_user, patient_columns=None):
    patient_columns = patient_columns or _patient_columns(cur)
    if "hospital_id" in patient_columns:
        cur.execute(
            """
            SELECT p.id
            FROM patients p
            LEFT JOIN lab_records r ON r.patient_id=p.id
            LEFT JOIN users u ON u.id=r.user_id
            WHERE p.id=%s AND (p.hospital_id=%s OR u.hospital_id=%s)
            LIMIT 1
            """,
            (patient_id, current_user["hospital_id"], current_user["hospital_id"]),
        )
    else:
        cur.execute(
            """
            SELECT p.id
            FROM patients p
            JOIN lab_records r ON r.patient_id=p.id
            JOIN users u ON u.id=r.user_id
            WHERE p.id=%s AND u.hospital_id=%s
            LIMIT 1
            """,
            (patient_id, current_user["hospital_id"]),
        )
    return cur.fetchone() is not None


def _missing_case_completion_items(cur, patient_id):
    cur.execute("SELECT name, gender, age, id_number FROM patients WHERE id=%s LIMIT 1", (patient_id,))
    patient = cur.fetchone() or {}
    missing = []
    if required(patient, ["name", "gender", "age", "id_number"]):
        missing.append("基础信息")

    required_tables = [
        ("diagnosis_records", "诊断"),
        ("assessments", "评估"),
        ("followups", "随访"),
        ("treatments", "治疗"),
    ]
    for table_name, label in required_tables:
        if not _table_exists(cur, table_name):
            missing.append(label)
            continue
        cur.execute(f"SELECT COUNT(*) AS total FROM `{table_name}` WHERE patient_id=%s", (patient_id,))
        count_row = cur.fetchone() or {}
        if int(count_row.get("total") or 0) < 1:
            missing.append(label)
    if not _required_lab_tests_completed(cur, patient_id):
        missing.append("检验")
    return missing


def _required_lab_tests_completed(cur, patient_id):
    if not _table_exists(cur, "lab_records"):
        return False
    cur.execute("SHOW COLUMNS FROM lab_records")
    columns = {row["Field"] for row in cur.fetchall()}
    select_parts = []
    if "lab_test_name" in columns:
        select_parts.append("lab_test_name")
    if "record_category" in columns:
        select_parts.append("record_category")
    if not select_parts:
        return False

    cur.execute(
        f"SELECT {','.join(select_parts)} FROM lab_records WHERE patient_id=%s ORDER BY created_at, id",
        (patient_id,),
    )
    names = []
    for row in cur.fetchall():
        lab_test_name = str(row.get("lab_test_name") or "").strip()
        record_category = str(row.get("record_category") or "").strip()
        category_label = (LAB_CATEGORY_DEFINITIONS.get(record_category) or {}).get("label", "")
        if lab_test_name:
            names.append(lab_test_name)
        if category_label and category_label != lab_test_name:
            names.append(category_label)
    if not names:
        return False

    prompt = """你是临床检验类别判断助手。请判断下面“已完成的检验类别”中，是否已经覆盖全部七类必需检验类别：血常规、生化电解质、血气分析、DIC7项、BNP、乳酸、PCT。
判断规则：
1. 名称可能带有编码或前缀，例如“JY2625.B型钠尿肽”，应理解为 BNP。
2. BNP 的同义名称包括：BNP、B型钠尿肽、N端B型钠尿肽、NT-proBNP、NT-pro BNP。
3. PCT 的同义名称包括：PCT、降钙素原。
4. 乳酸的同义名称包括：乳酸、LAC、LA。
5. DIC7项可写作：DIC7项、DIC 7项、凝血、凝血功能、凝血七项。
6. 只有七类都已覆盖，才返回“是”；只要缺少任意一类，就返回“否”。
7. 只能返回一个字：是 或 否，不要解释。

已完成的检验名称：
""" + "\n".join(f"- {name}" for name in names)
    try:
        return ask_ai_yes_no(prompt, timeout=60)
    except Exception as e:
        print(f"Lab completion AI check error: {e}")
        return False


def _mark_case_complete_if_ready(cur, patient_id, patient_columns=None):
    patient_columns = patient_columns or _patient_columns(cur)
    if "case_status" not in patient_columns and "case_integrity" not in patient_columns:
        return False

    status_fields = []
    if "case_status" in patient_columns:
        status_fields.append("case_status")
    if "case_integrity" in patient_columns:
        status_fields.append("case_integrity")
    cur.execute(
        f"SELECT {','.join(status_fields)} FROM patients WHERE id=%s LIMIT 1",
        (patient_id,),
    )
    patient = cur.fetchone() or {}
    if patient.get("case_status") == "submitted" or patient.get("case_integrity") == "complete":
        return False
    if _missing_case_completion_items(cur, patient_id):
        return False

    updates = []
    if "case_status" in patient_columns:
        updates.append("case_status='submitted'")
    if "case_integrity" in patient_columns:
        updates.append("case_integrity='complete'")
    if "updated_at" in patient_columns:
        updates.append("updated_at=NOW()")
    cur.execute(f"UPDATE patients SET {','.join(updates)} WHERE id=%s", (patient_id,))
    return True


def _find_existing_patient_for_lab(cur, current_user, patient, patient_columns):
    conditions = ["name=%s", "IFNULL(gender,'')=%s", "IFNULL(age,'')=%s"]
    params = [patient.get("name"), patient.get("gender", ""), patient.get("age") or ""]
    if "hospital_id" in patient_columns:
        conditions.insert(0, "hospital_id=%s")
        params.insert(0, current_user["hospital_id"])
    cur.execute(
        f"SELECT id FROM patients WHERE {' AND '.join(conditions)} ORDER BY id LIMIT 1",
        params,
    )
    row = cur.fetchone()
    return row["id"] if row else None


def _find_duplicate_lab_record(cur, patient_id, existing_columns, record_category, lab_test_name, photo_path, record_values):
    if photo_path and "photo_path" in existing_columns:
        cur.execute(
            "SELECT id FROM lab_records WHERE patient_id=%s AND photo_path=%s LIMIT 1",
            (patient_id, photo_path),
        )
        row = cur.fetchone()
        if row:
            return row["id"]

    conditions = ["patient_id=%s"]
    params = [patient_id]
    if "record_category" in existing_columns:
        conditions.append("IFNULL(record_category,'')=%s")
        params.append(record_category or "")
    if "lab_test_name" in existing_columns:
        conditions.append("IFNULL(lab_test_name,'')=%s")
        params.append(lab_test_name or "")
    for field, value in record_values.items():
        conditions.append(f"IFNULL(`{field.replace('`', '``')}`,'')=%s")
        params.append(value or "")
    cur.execute(
        f"SELECT id FROM lab_records WHERE {' AND '.join(conditions)} ORDER BY id LIMIT 1",
        params,
    )
    row = cur.fetchone()
    return row["id"] if row else None


@records_bp.get("/api/diseases")
@require_user
def diseases():
    conn = db()
    try:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT d.id, d.name, COUNT(DISTINCT r.patient_id) AS patient_count
                FROM diseases d
                LEFT JOIN lab_records r ON r.disease_id=d.id
                GROUP BY d.id, d.name, d.sort_order
                ORDER BY d.sort_order, d.id
                """
            )
            diseases = cur.fetchall()

            cur.execute("SELECT COUNT(DISTINCT patient_id) AS total_patients FROM lab_records")
            total_row = cur.fetchone() or {}
            total_patients = total_row.get("total_patients", 0) or 0

            return ok({"diseases": diseases, "total_patients": int(total_patients)})
    finally:
        conn.close()


@records_bp.get("/api/member-reviews")
@require_user
def member_reviews():
    user_id = session.get("user_id")
    conn = db()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT id, hospital_id, parent_id FROM users WHERE id=%s", (user_id,))
            current_user = cur.fetchone()
            if not current_user:
                return fail("用户不存在", 404)
            if int(current_user.get("parent_id", 0)) != 0:
                return fail("无权限访问会员审核", 403)

            cur.execute(
                """
                SELECT id, name, phone, status
                FROM users
                WHERE hospital_id=%s AND id<>%s
                ORDER BY id DESC
                """,
                (current_user["hospital_id"], user_id),
            )
            return ok(cur.fetchall())
    finally:
        conn.close()


@records_bp.post("/api/member-reviews/<int:target_user_id>/approve")
@require_user
def approve_member(target_user_id):
    user_id = session.get("user_id")
    conn = db()
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT id, hospital_id, parent_id FROM users WHERE id=%s", (user_id,))
            current_user = cur.fetchone()
            if not current_user:
                return fail("用户不存在", 404)
            if int(current_user.get("parent_id", 0)) != 0:
                return fail("无权限操作", 403)

            cur.execute("SELECT id, hospital_id, status FROM users WHERE id=%s", (target_user_id,))
            target_user = cur.fetchone()
            if not target_user:
                return fail("目标会员不存在", 404)
            if target_user["hospital_id"] != current_user["hospital_id"]:
                return fail("仅可审核同医院成员", 403)

            if target_user["status"] != "pending":
                return fail("该会员不是未审核状态")

            cur.execute("UPDATE users SET status='approved' WHERE id=%s", (target_user_id,))
        conn.commit()
        return ok(message="审核通过")
    finally:
        conn.close()


@records_bp.get("/api/forms/<int:disease_id>")
@require_user
def form_fields(disease_id):
    return ok(get_fields_for_disease(disease_id))


@records_bp.post("/api/recognize")
@require_user
def recognize():
    raw_id = request.form.get("disease_id", "0")
    recognize_mode = (request.form.get("mode") or "lab").strip()
    record_category = (request.form.get("record_category") or "").strip()
    try:
        disease_id = int(raw_id)
    except (ValueError, TypeError):
        disease_id = 0

    photo = request.files.get("photo")
    photo_path = None
    if photo and photo.filename:
        os.makedirs(current_app.config["UPLOAD_FOLDER"], exist_ok=True)
        filename = f"{int(time.time())}_{secure_filename(photo.filename)}"
        path = os.path.join(current_app.config["UPLOAD_FOLDER"], filename)
        photo.save(path)
        photo_path = path.replace("\\", "/")

    fields = get_fields_for_disease(disease_id)
    if not fields:
        conn = db()
        try:
            with conn.cursor() as cur:
                cur.execute("SELECT id, field_name, data_type, comment_text, form_label FROM field_settings WHERE enabled=1 ORDER BY id")
                fields = cur.fetchall()
        finally:
            conn.close()

    values = {}
    patient = {}
    lab_test_name = ""
    prompt_text = INTAKE_PROMPT if recognize_mode == "intake" else None
    if recognize_mode == "lab" and record_category:
        category_config = _lab_category_config(record_category)
        if category_config:
            prompt_text = _build_category_prompt(category_config)
    if photo_path:
        try:
            ai_text = recognize_image(photo_path, prompt_text=prompt_text)
            if not ai_text:
                return fail("AI识别接口暂无返回，请稍后重试", 502)
            values = {} if recognize_mode == "intake" else parse_ai_result(ai_text, fields)
            patient = parse_patient_info(ai_text)
            lab_test_name = parse_lab_test_name(ai_text)
            fields = reorder_fields_by_values(fields, values)
        except Exception as e:
            print(f"AI recognition error: {e}")

    return ok(
        {
            "photo_path": photo_path,
            "fields": fields,
            "values": values,
            "patient": patient,
            "mode": recognize_mode,
            "record_category": record_category or None,
            "category_label": (_lab_category_config(record_category) or {}).get("label") if record_category else None,
            "lab_test_name": lab_test_name,
        },
        "识别完成",
    )


@records_bp.post("/api/records")
@require_user
def save_record():
    data = request.get_json(silent=True) or {}
    patient_id = data.get("patient_id")
    raw_patient = data.get("patient") or {}
    patient = {
        "name": (raw_patient.get("name") or "").strip(),
        "gender": (raw_patient.get("gender") or "").strip(),
        "age": raw_patient.get("age"),
        "phone": (raw_patient.get("phone") or "").strip(),
        "id_number": (raw_patient.get("id_number") or "").strip(),
    }
    values = data.get("values") or {}
    disease_id = data.get("disease_id")
    record_category = (data.get("record_category") or "").strip() or None
    lab_test_name = (data.get("lab_test_name") or "").strip() or None
    if not disease_id:
        return fail("请先选择疾病")
    if not patient_id and required(patient, ["name"]):
        return fail("请填写病患姓名")

    allowed_fields = {field["field_name"] for field in get_fields_for_disease(int(disease_id))}

    conn = db()
    try:
        with conn.cursor() as cur:
            current_user = _current_user_scope(cur)
            if not current_user:
                return fail("用户不存在", 404)

            cur.execute("SHOW COLUMNS FROM lab_records")
            existing_columns = {row["Field"] for row in cur.fetchall()}
            patient_columns = _patient_columns(cur)

            # Only keep fields that are configured for the disease and already exist
            # in lab_records to avoid runtime SQL errors from stale field settings.
            skipped_fields = sorted(
                [key for key in values.keys() if key in allowed_fields and key not in existing_columns]
            )

            record_values = {
                key: value
                for key, value in values.items()
                if key in allowed_fields and key in existing_columns
            }

            if patient_id:
                patient_id = int(patient_id)
                if not _patient_is_accessible(cur, patient_id, current_user, patient_columns):
                    return fail("病人不存在", 404)
                updates = []
                params = []
                if patient.get("name"):
                    updates.append("name=%s")
                    params.append(patient.get("name"))
                if patient.get("gender"):
                    updates.append("gender=%s")
                    params.append(patient.get("gender"))
                if patient.get("age") not in (None, ""):
                    updates.append("age=%s")
                    params.append(patient.get("age") or None)
                if patient.get("phone") is not None:
                    updates.append("phone=%s")
                    params.append(patient.get("phone"))
                if patient.get("id_number"):
                    updates.append("id_number=%s")
                    params.append(patient.get("id_number"))
                if "last_disease_id" in patient_columns:
                    updates.append("last_disease_id=%s")
                    params.append(disease_id)
                if "case_status" in patient_columns:
                    updates.append("case_status='draft'")
                updates.append("updated_at=NOW()")
                params.append(patient_id)
                cur.execute(f"UPDATE patients SET {','.join(updates)} WHERE id=%s", params)
            else:
                existing_patient_id = _find_existing_patient_for_lab(cur, current_user, patient, patient_columns)
                if existing_patient_id:
                    patient_id = existing_patient_id
                else:
                    patient_insert = {
                        "name": patient.get("name"),
                        "gender": patient.get("gender"),
                        "age": patient.get("age") or None,
                        "phone": patient.get("phone"),
                        "id_number": patient.get("id_number"),
                    }
                    if "hospital_id" in patient_columns:
                        patient_insert["hospital_id"] = current_user["hospital_id"]
                    if "user_id" in patient_columns:
                        patient_insert["user_id"] = current_user["id"]
                    if "case_status" in patient_columns:
                        patient_insert["case_status"] = "draft"
                    if "last_disease_id" in patient_columns:
                        patient_insert["last_disease_id"] = disease_id
                    patient_cols = list(patient_insert.keys())
                    patient_placeholders = ",".join(["%s"] * len(patient_cols))
                    cur.execute(
                        f"INSERT INTO patients ({','.join(patient_cols)}) VALUES ({patient_placeholders})",
                        list(patient_insert.values()),
                    )
                    patient_id = cur.lastrowid

            duplicate_record_id = _find_duplicate_lab_record(
                cur,
                patient_id,
                existing_columns,
                record_category,
                lab_test_name,
                data.get("photo_path"),
                record_values,
            )
            if duplicate_record_id:
                return ok({"record_id": duplicate_record_id, "patient_id": patient_id, "duplicate": True}, "本次检验指标已存在，未重复保存")

            columns = ["patient_id", "user_id", "disease_id", "photo_path", "ai_raw"]
            params = [patient_id, session["user_id"], disease_id, data.get("photo_path"), data.get("ai_raw")]
            if "record_category" in existing_columns:
                columns.append("record_category")
                params.append(record_category)
            if "lab_test_name" in existing_columns:
                columns.append("lab_test_name")
                params.append(lab_test_name)
            columns.extend(list(record_values.keys()))
            params.extend(list(record_values.values()))
            placeholders = ",".join(["%s"] * len(columns))
            # PyMySQL uses Python '%' formatting internally; escape '%' in identifiers
            # to avoid '%`' formatting errors for legacy field names like 'xxx%'.
            safe_columns = ",".join([f"`{str(col).replace('`', '``').replace('%', '%%')}`" for col in columns])
            cur.execute(f"INSERT INTO lab_records ({safe_columns}) VALUES ({placeholders})", params)
            record_id = cur.lastrowid
            completed_now = _mark_case_complete_if_ready(cur, patient_id, patient_columns)
        conn.commit()
        message = "保存成功，已自动转为完整记录" if completed_now else "保存成功"
        if skipped_fields:
            return ok(
                {"record_id": record_id, "patient_id": patient_id, "skipped_fields": skipped_fields, "completed_now": completed_now},
                message + "，部分字段未入库",
            )
        return ok({"record_id": record_id, "patient_id": patient_id, "completed_now": completed_now}, message)
    finally:
        conn.close()


@records_bp.get("/api/records/<int:record_id>")
@require_user
def record_detail(record_id):
    conn = db()
    try:
        with conn.cursor() as cur:
            current_user = _current_user_scope(cur)
            if not current_user:
                return fail("用户不存在", 404)

            patient_columns = _patient_columns(cur)
            if "hospital_id" in patient_columns:
                cur.execute(
                    """
                    SELECT r.*, p.name AS patient_name, p.gender, p.age, p.phone, p.id_number,
                           d.name AS disease_name, u.name AS operator_name
                    FROM lab_records r
                    JOIN patients p ON p.id=r.patient_id
                    LEFT JOIN diseases d ON d.id=r.disease_id
                    LEFT JOIN users u ON u.id=r.user_id
                    WHERE r.id=%s AND (p.hospital_id=%s OR u.hospital_id=%s)
                    LIMIT 1
                    """,
                    (record_id, current_user["hospital_id"], current_user["hospital_id"]),
                )
            else:
                cur.execute(
                    """
                    SELECT r.*, p.name AS patient_name, p.gender, p.age, p.phone, p.id_number,
                           d.name AS disease_name, u.name AS operator_name
                    FROM lab_records r
                    JOIN patients p ON p.id=r.patient_id
                    JOIN users u ON u.id=r.user_id
                    LEFT JOIN diseases d ON d.id=r.disease_id
                    WHERE r.id=%s AND u.hospital_id=%s
                    LIMIT 1
                    """,
                    (record_id, current_user["hospital_id"]),
                )
            record = cur.fetchone()
            if not record:
                return fail("检验记录不存在", 404)

            cur.execute("SHOW COLUMNS FROM lab_records")
            record_columns = {row["Field"] for row in cur.fetchall()}
            cur.execute(
                """
                SELECT field_name, form_label, unit, reference_range, test_method
                FROM field_settings
                WHERE enabled=1
                ORDER BY id
                """
            )
            system_fields = {"id", "patient_id", "user_id", "disease_id", "record_category", "lab_test_name", "photo_path", "ai_raw", "created_at"}
            items = []
            for field in cur.fetchall():
                field_name = field["field_name"]
                if field_name not in record_columns or field_name in system_fields:
                    continue
                value = record.get(field_name)
                if value is None or str(value).strip() == "":
                    continue
                items.append({
                    "field_name": field_name,
                    "form_label": field["form_label"],
                    "value": value,
                    "unit": field.get("unit"),
                    "reference_range": field.get("reference_range"),
                    "test_method": field.get("test_method"),
                })

            return ok({"record": record, "items": items})
    finally:
        conn.close()


@records_bp.get("/api/cases/export")
@require_user
def export_cases():
    conn = db()
    try:
        with conn.cursor() as cur:
            current_user = _current_user_scope(cur)
            if not current_user:
                return fail("用户不存在", 404)

            multiplier = 1
            if _table_exists(cur, "system_settings"):
                cur.execute("SELECT doctor_download_multiplier FROM system_settings LIMIT 1")
                settings = cur.fetchone() or {}
                try:
                    multiplier = max(int(settings.get("doctor_download_multiplier", 1)), 0)
                except (TypeError, ValueError):
                    multiplier = 1

            patient_columns = _patient_columns(cur)
            own_conditions = ["r.user_id=%s"]
            own_params = [current_user["id"]]
            if "user_id" in patient_columns:
                own_conditions.append("p.user_id=%s")
                own_params.append(current_user["id"])

            cur.execute(
                f"""
                SELECT DISTINCT p.id
                FROM patients p
                LEFT JOIN lab_records r ON r.patient_id=p.id
                WHERE {' OR '.join(own_conditions)}
                """,
                own_params,
            )
            own_patient_ids = [row["id"] for row in cur.fetchall()]
            other_limit = len(own_patient_ids) * multiplier

            other_patient_ids = []
            if other_limit > 0:
                exclude_placeholders = ",".join(["%s"] * len(own_patient_ids))
                exclude_sql = f"AND p.id NOT IN ({exclude_placeholders})" if own_patient_ids else ""
                other_conditions = ["r.user_id IS NOT NULL AND r.user_id<>%s"]
                other_params = [current_user["id"]]
                if "user_id" in patient_columns:
                    other_conditions.append("p.user_id IS NOT NULL AND p.user_id<>%s")
                    other_params.append(current_user["id"])
                cur.execute(
                    f"""
                    SELECT DISTINCT p.id
                    FROM patients p
                    LEFT JOIN lab_records r ON r.patient_id=p.id
                    WHERE ({' OR '.join(other_conditions)}) {exclude_sql}
                    ORDER BY RAND()
                    LIMIT {int(other_limit)}
                    """,
                    other_params + own_patient_ids,
                )
                other_patient_ids = [row["id"] for row in cur.fetchall()]

            patient_ids = own_patient_ids + other_patient_ids
            patient_rows, patient_headers = _fetch_rows_for_patients(cur, "patients", patient_ids, "id")
            lab_rows, lab_headers = _fetch_rows_for_patients(cur, "lab_records", patient_ids)
            treatment_rows, treatment_headers = _fetch_rows_for_patients(cur, "treatments", patient_ids)
            followup_rows, followup_headers = _fetch_rows_for_patients(cur, "followups", patient_ids)

            workbook = Workbook()
            workbook.remove(workbook.active)
            _write_sheet(workbook, "病患信息", patient_rows, patient_headers)
            _write_sheet(workbook, "检验记录", lab_rows, lab_headers)
            _write_sheet(workbook, "治疗记录", treatment_rows, treatment_headers)
            _write_sheet(workbook, "随访记录", followup_rows, followup_headers)

            output = BytesIO()
            workbook.save(output)
            output.seek(0)
            filename = f"case_export_{int(time.time())}.xlsx"
            return send_file(
                output,
                as_attachment=True,
                download_name=filename,
                mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
    finally:
        conn.close()


@records_bp.get("/api/cases")
@require_user
def cases():
    conn = db()
    try:
        with conn.cursor() as cur:
            current_user = _current_user_scope(cur)
            if not current_user:
                return fail("用户不存在", 404)

            patient_columns = _patient_columns(cur)
            status_expr = "IFNULL(p.case_status, 'draft')" if "case_status" in patient_columns else "'draft'"
            integrity_expr = "IFNULL(p.case_integrity, IFNULL(p.case_status, 'draft'))" if "case_integrity" in patient_columns and "case_status" in patient_columns else ("IFNULL(p.case_integrity, 'draft')" if "case_integrity" in patient_columns else status_expr)
            disease_join = "LEFT JOIN diseases d ON d.id=p.last_disease_id" if "last_disease_id" in patient_columns else ""
            disease_select = ", d.name AS disease_name" if "last_disease_id" in patient_columns else ", NULL AS disease_name"

            if "hospital_id" in patient_columns:
                cur.execute(
                    """
                    SELECT p.id, p.name, p.gender, p.age, p.phone, p.id_number,
                           COUNT(r.id) AS record_count,
                           MAX(r.created_at) AS latest_record_at,
                           {status_expr} AS case_status,
                           {integrity_expr} AS case_integrity
                           {disease_select}
                    FROM patients p
                    LEFT JOIN lab_records r ON r.patient_id=p.id
                    LEFT JOIN users u ON u.id=r.user_id
                    {disease_join}
                    WHERE p.hospital_id=%s OR u.hospital_id=%s
                    GROUP BY p.id, p.name, p.gender, p.age, p.phone, p.id_number, p.created_at, case_status, case_integrity, disease_name
                    ORDER BY IFNULL(MAX(r.created_at), p.created_at) DESC, p.id DESC
                    """.format(status_expr=status_expr, integrity_expr=integrity_expr, disease_select=disease_select, disease_join=disease_join),
                    (current_user["hospital_id"], current_user["hospital_id"]),
                )
            else:
                cur.execute(
                    """
                    SELECT p.id, p.name, p.gender, p.age, p.phone, p.id_number,
                           COUNT(r.id) AS record_count,
                           MAX(r.created_at) AS latest_record_at,
                           {status_expr} AS case_status,
                           {integrity_expr} AS case_integrity
                           {disease_select}
                    FROM patients p
                    JOIN lab_records r ON r.patient_id=p.id
                    JOIN users u ON u.id=r.user_id
                    {disease_join}
                    WHERE u.hospital_id=%s
                    GROUP BY p.id, p.name, p.gender, p.age, p.phone, p.id_number, case_status, case_integrity, disease_name
                    ORDER BY latest_record_at DESC, p.id DESC
                    """.format(status_expr=status_expr, integrity_expr=integrity_expr, disease_select=disease_select, disease_join=disease_join),
                    (current_user["hospital_id"],),
                )
            rows = cur.fetchall()
            return ok({
                "draft": [row for row in rows if row.get("case_integrity") not in {"complete", "submitted"}],
                "complete": [row for row in rows if row.get("case_integrity") in {"complete", "submitted"}],
            })
    finally:
        conn.close()


@records_bp.post("/api/patients")
@require_user
def create_patient():
    data = request.get_json(silent=True) or request.form.to_dict()
    require_age = str(data.get("require_age") or "1") != "0"
    patient = {
        "name": (data.get("name") or "").strip(),
        "gender": (data.get("gender") or "").strip(),
        "age": data.get("age") or None,
        "phone": (data.get("phone") or "").strip(),
        "id_number": (data.get("id_number") or "").strip(),
    }
    required_fields = ["name", "gender", "id_number"]
    if require_age:
        required_fields.append("age")
    if required(patient, required_fields):
        return fail("请填写姓名、性别、登记号" + ("、年龄" if require_age else ""))

    conn = db()
    try:
        with conn.cursor() as cur:
            current_user = _current_user_scope(cur)
            if not current_user:
                return fail("用户不存在", 404)

            patient_columns = _patient_columns(cur)
            if "hospital_id" not in patient_columns:
                return fail("请先执行 ensure_patient_case_owner_columns.sql 后再新增病例")
            if "hospital_id" in patient_columns:
                cur.execute(
                    """
                    SELECT id FROM patients
                    WHERE hospital_id=%s AND name=%s AND IFNULL(phone,'')=%s AND IFNULL(id_number,'')=%s
                    LIMIT 1
                    """,
                    (current_user["hospital_id"], patient["name"], patient["phone"], patient["id_number"]),
                )
            else:
                cur.execute(
                    """
                    SELECT id FROM patients
                    WHERE name=%s AND IFNULL(phone,'')=%s AND IFNULL(id_number,'')=%s
                    LIMIT 1
                    """,
                    (patient["name"], patient["phone"], patient["id_number"]),
                )
            existing = cur.fetchone()
            last_disease_id = data.get("last_disease_id") or None
            case_status = (data.get("case_status") or "draft").strip() or "draft"
            case_integrity = (data.get("case_integrity") or "draft").strip() or "draft"
            if existing:
                updates = []
                params = []
                if "last_disease_id" in patient_columns and last_disease_id:
                    updates.append("last_disease_id=%s")
                    params.append(last_disease_id)
                if "case_status" in patient_columns:
                    updates.append("case_status=%s")
                    params.append(case_status)
                if "case_integrity" in patient_columns:
                    updates.append("case_integrity=%s")
                    params.append(case_integrity)
                updates.append("updated_at=NOW()")
                params.append(existing["id"])
                cur.execute(f"UPDATE patients SET {','.join(updates)} WHERE id=%s", params)
                return ok({"patient_id": existing["id"]}, "病例已存在")

            patient_insert = dict(patient)
            if "hospital_id" in patient_columns:
                patient_insert["hospital_id"] = current_user["hospital_id"]
            if "user_id" in patient_columns:
                patient_insert["user_id"] = current_user["id"]
            if "case_status" in patient_columns:
                patient_insert["case_status"] = case_status
            if "case_integrity" in patient_columns:
                patient_insert["case_integrity"] = case_integrity
            if "last_disease_id" in patient_columns and last_disease_id:
                patient_insert["last_disease_id"] = last_disease_id

            columns = list(patient_insert.keys())
            placeholders = ",".join(["%s"] * len(columns))
            cur.execute(
                f"INSERT INTO patients ({','.join(columns)}) VALUES ({placeholders})",
                list(patient_insert.values()),
            )
            patient_id = cur.lastrowid
        conn.commit()
        return ok({"patient_id": patient_id}, "病例已新增")
    finally:
        conn.close()


@records_bp.get("/api/patients/<int:patient_id>")
@require_user
def patient_detail(patient_id):
    conn = db()
    try:
        with conn.cursor() as cur:
            current_user = _current_user_scope(cur)
            if not current_user:
                return fail("用户不存在", 404)

            patient_columns = _patient_columns(cur)
            if "hospital_id" in patient_columns:
                cur.execute(
                    """
                    SELECT p.*
                    FROM patients p
                    LEFT JOIN lab_records r ON r.patient_id=p.id
                    LEFT JOIN users u ON u.id=r.user_id
                    WHERE p.id=%s AND (p.hospital_id=%s OR u.hospital_id=%s)
                    LIMIT 1
                    """,
                    (patient_id, current_user["hospital_id"], current_user["hospital_id"]),
                )
            else:
                cur.execute(
                    """
                    SELECT p.*
                    FROM patients p
                    JOIN lab_records r ON r.patient_id=p.id
                    JOIN users u ON u.id=r.user_id
                    WHERE p.id=%s AND u.hospital_id=%s
                    LIMIT 1
                    """,
                    (patient_id, current_user["hospital_id"]),
                )
            patient = cur.fetchone()
            if not patient:
                return fail("病人不存在", 404)

            cur.execute(
                """
                SELECT r.id, r.created_at, r.record_category, r.lab_test_name, d.name AS disease_name, u.name AS operator_name
                FROM lab_records r
                LEFT JOIN diseases d ON d.id=r.disease_id
                LEFT JOIN users u ON u.id=r.user_id
                WHERE r.patient_id=%s
                ORDER BY r.created_at DESC, r.id DESC
                """,
                (patient_id,),
            )
            lab_records = cur.fetchall()

            cur.execute("SHOW COLUMNS FROM treatments")
            treatment_columns = {row["Field"] for row in cur.fetchall()}
            if "diagnosis_record_id" in treatment_columns and _table_exists(cur, "diagnosis_records"):
                cur.execute(
                    """
                    SELECT t.*, dr.diagnosis_disease, dr.preliminary_diagnosis
                    FROM treatments t
                    LEFT JOIN diagnosis_records dr ON dr.id=t.diagnosis_record_id
                    WHERE t.patient_id=%s
                    ORDER BY t.treat_time DESC, t.id DESC
                    """,
                    (patient_id,),
                )
            else:
                cur.execute(
                    """
                    SELECT *
                    FROM treatments
                    WHERE patient_id=%s
                    ORDER BY treat_time DESC, id DESC
                    """,
                    (patient_id,),
                )
            treatments = cur.fetchall()

            cur.execute(
                """
                SELECT f.*, dr.diagnosis_disease, dr.preliminary_diagnosis
                FROM followups f
                LEFT JOIN diagnosis_records dr ON dr.id=f.diagnosis_record_id
                WHERE f.patient_id=%s
                ORDER BY f.follow_time DESC, f.id DESC
                """,
                (patient_id,),
            )
            followups = cur.fetchall()

            assessments = []
            if _table_exists(cur, "assessments"):
                cur.execute(
                    """
                    SELECT a.*, dr.diagnosis_disease, dr.preliminary_diagnosis
                    FROM assessments a
                    LEFT JOIN diagnosis_records dr ON dr.id=a.diagnosis_record_id
                    WHERE a.patient_id=%s
                    ORDER BY a.assessment_time DESC, a.id DESC
                    """,
                    (patient_id,),
                )
                assessments = cur.fetchall()

            diagnosis_records = []
            if _table_exists(cur, "diagnosis_records"):
                cur.execute(
                    """
                    SELECT dr.id, dr.diagnosis_time, dr.diagnosis_disease,
                           dr.preliminary_diagnosis, dr.medical_history, u.name AS operator_name
                    FROM diagnosis_records dr
                    LEFT JOIN users u ON u.id=dr.user_id
                    WHERE dr.patient_id=%s
                    ORDER BY dr.diagnosis_time DESC, dr.id DESC
                    """,
                    (patient_id,),
                )
                diagnosis_records = cur.fetchall()

            base_fields = {
                "id", "hospital_id", "user_id", "name", "gender", "age", "phone", "id_number",
                "diagnosis", "diagnosis_disease", "medical_history", "preliminary_diagnosis",
                "case_integrity", "created_at", "updated_at",
            }
            cur.execute("SELECT field_name, form_label FROM patient_field_settings WHERE enabled=1 ORDER BY created_at DESC")
            patient_meta = []
            for field in cur.fetchall():
                field_name = field["field_name"]
                if field_name in patient_columns and field_name not in base_fields:
                    patient_meta.append({
                        "field_name": field_name,
                        "form_label": field["form_label"],
                        "value": patient.get(field_name),
                    })

            return ok({
                "patient": patient,
                "lab_records": lab_records,
                "treatments": treatments,
                "followups": followups,
                "assessments": assessments,
                "diagnosis_records": diagnosis_records,
                "patient_fields": patient_meta,
                "lab_categories": [
                    {"key": key, "label": value["label"]}
                    for key, value in LAB_CATEGORY_DEFINITIONS.items()
                ],
            })
    finally:
        conn.close()


@records_bp.post("/api/patients/<int:patient_id>")
@require_user
def update_patient(patient_id):
    data = request.get_json(silent=True) or request.form.to_dict()
    is_diagnosis_update = any(
        field in data for field in {"diagnosis_disease", "medical_history", "preliminary_diagnosis"}
    )
    if not is_diagnosis_update:
        required_fields = ["name", "gender", "age", "id_number"]
        if required(data, required_fields):
            return fail("请填写姓名、性别、年龄、病历号")

    conn = db()
    try:
        with conn.cursor() as cur:
            current_user = _current_user_scope(cur)
            if not current_user:
                return fail("用户不存在", 404)

            patient_columns = _patient_columns(cur)
            if not _patient_is_accessible(cur, patient_id, current_user, patient_columns):
                return fail("病人不存在", 404)

            allowed_fields = {"name", "gender", "age", "phone", "id_number"}
            cur.execute("SELECT field_name FROM patient_field_settings WHERE enabled=1")
            allowed_fields.update(
                row["field_name"] for row in cur.fetchall() if row["field_name"] in patient_columns
            )
            allowed_fields = [field for field in allowed_fields if field in patient_columns]

            if "diagnosis_disease" in data:
                diagnosis_disease = (data.get("diagnosis_disease") or "").strip()
                if is_diagnosis_update and not diagnosis_disease:
                    return fail("请选择本次诊断疾病")
                if diagnosis_disease and diagnosis_disease not in DIAGNOSIS_DISEASE_OPTIONS:
                    return fail("疾病选项不合法")
                data["diagnosis_disease"] = diagnosis_disease
            if "medical_history" in data:
                data["medical_history"] = _normalize_medical_history(data.get("medical_history"))
            if "preliminary_diagnosis" in data:
                data["preliminary_diagnosis"] = _normalize_preliminary_diagnosis(
                    data.get("diagnosis_disease", ""),
                    data.get("preliminary_diagnosis"),
                )

            updates = []
            params = []
            for field in allowed_fields:
                if field not in data:
                    continue
                value = data.get(field)
                if field == "age":
                    value = value or None
                updates.append(f"`{field.replace('`', '``')}`=%s")
                params.append(value)

            if not updates:
                return fail("没有可保存的基础信息")
            if "updated_at" in patient_columns:
                updates.append("updated_at=NOW()")
            params.append(patient_id)
            cur.execute(f"UPDATE patients SET {','.join(updates)} WHERE id=%s", params)

            diagnosis_disease = (data.get("diagnosis_disease") or "").strip()
            if is_diagnosis_update and diagnosis_disease and _table_exists(cur, "diagnosis_records"):
                cur.execute(
                    """
                    INSERT INTO diagnosis_records
                      (patient_id, user_id, diagnosis_time, diagnosis_disease, preliminary_diagnosis, medical_history)
                    VALUES (%s,%s,NOW(),%s,%s,%s)
                    """,
                    (
                        patient_id,
                        current_user["id"],
                        diagnosis_disease,
                        data.get("preliminary_diagnosis") or None,
                        data.get("medical_history") or "无",
                    ),
                )
            completed_now = _mark_case_complete_if_ready(cur, patient_id, patient_columns)
        conn.commit()
        message = "诊断信息已保存" if is_diagnosis_update else "基础信息已保存"
        if completed_now:
            message += "，已自动转为完整记录"
        return ok({"patient_id": patient_id, "completed_now": completed_now}, message)
    finally:
        conn.close()


@records_bp.post("/api/patients/<int:patient_id>/treatments")
@require_user
def add_treatment(patient_id):
    payload = request.get_json(silent=True) or request.form.to_dict()
    treat_time = (payload.get("treat_time") or "").strip()
    diagnosis_record_id = payload.get("diagnosis_record_id")
    if not diagnosis_record_id:
        return fail("请选择诊断记录")
    if not treat_time:
        treat_time = datetime.now().strftime("%Y-%m-%dT%H:%M")

    conn = db()
    try:
        with conn.cursor() as cur:
            current_user = _current_user_scope(cur)
            if not current_user:
                return fail("用户不存在", 404)
            if not _patient_is_accessible(cur, patient_id, current_user):
                return fail("病人不存在", 404)
            if not _table_exists(cur, "diagnosis_records"):
                return fail("诊断记录表不存在，请先执行 ensure_diagnosis_records.sql")

            cur.execute(
                "SELECT diagnosis_disease FROM diagnosis_records WHERE id=%s AND patient_id=%s LIMIT 1",
                (diagnosis_record_id, patient_id),
            )
            diagnosis = cur.fetchone()
            if not diagnosis:
                return fail("请选择有效的诊断记录")

            cur.execute("SHOW COLUMNS FROM treatments")
            treatment_columns = {row["Field"] for row in cur.fetchall()}
            values = {
                "patient_id": patient_id,
                "user_id": session.get("user_id"),
                "diagnosis_record_id": diagnosis_record_id,
                "treat_time": treat_time,
                "treatment_method": "结构化治疗记录",
                "detail_json": json.dumps(payload, ensure_ascii=False),
                "antibiotics": (payload.get("antibiotics") or "").strip() or None,
                "antibiotics_start_time": _optional_datetime(payload.get("antibiotics_start_time")),
                "vasoactive_drugs": (payload.get("vasoactive_drugs") or "").strip() or None,
                "vasoactive_start_time": _optional_datetime(payload.get("vasoactive_start_time")),
                "vasoactive_concentration": (payload.get("vasoactive_concentration") or "").strip() or None,
                "volume_management": (payload.get("volume_management") or "").strip() or None,
                "volume_total_ml": _optional_decimal(payload.get("volume_total_ml")),
                "respiratory_support": (payload.get("respiratory_support") or "").strip() or None,
                "respiratory_start_time": _optional_datetime(payload.get("respiratory_start_time")),
                "immunomodulators": (payload.get("immunomodulators") or "").strip() or None,
                "immunomodulator_start_time": _optional_datetime(payload.get("immunomodulator_start_time")),
                "blood_purification": (payload.get("blood_purification") or "").strip() or None,
                "blood_purification_start_time": _optional_datetime(payload.get("blood_purification_start_time")),
                "traditional_chinese_medicine": (payload.get("traditional_chinese_medicine") or "").strip() or None,
                "traditional_chinese_medicine_start_time": _optional_datetime(payload.get("traditional_chinese_medicine_start_time")),
                "digestive_secretion_drugs": (payload.get("digestive_secretion_drugs") or "").strip() or None,
                "digestive_secretion_drugs_start_time": _optional_datetime(payload.get("digestive_secretion_drugs_start_time")),
                "cardiac_treatment_methods": (payload.get("cardiac_treatment_methods") or "").strip() or None,
                "cardiac_treatment_start_time": _optional_datetime(payload.get("cardiac_treatment_start_time")),
                "sodium_channel_blockers": (payload.get("sodium_channel_blockers") or "").strip() or None,
                "sodium_channel_blocker_start_time": _optional_datetime(payload.get("sodium_channel_blocker_start_time")),
                "beta_blockers": (payload.get("beta_blockers") or "").strip() or None,
                "beta_blocker_start_time": _optional_datetime(payload.get("beta_blocker_start_time")),
                "potassium_channel_blockers": (payload.get("potassium_channel_blockers") or "").strip() or None,
                "potassium_channel_blocker_start_time": _optional_datetime(payload.get("potassium_channel_blocker_start_time")),
                "calcium_channel_blockers": (payload.get("calcium_channel_blockers") or "").strip() or None,
                "calcium_channel_blocker_start_time": _optional_datetime(payload.get("calcium_channel_blocker_start_time")),
                "other_cardiac_drugs": (payload.get("other_cardiac_drugs") or "").strip() or None,
                "other_cardiac_drugs_start_time": _optional_datetime(payload.get("other_cardiac_drugs_start_time")),
                "poisoning_other_drugs": (payload.get("poisoning_other_drugs") or "").strip() or None,
                "poisoning_other_drugs_start_time": _optional_datetime(payload.get("poisoning_other_drugs_start_time")),
                "intracranial_pressure_reduction": (payload.get("intracranial_pressure_reduction") or "").strip() or None,
                "intracranial_pressure_start_time": _optional_datetime(payload.get("intracranial_pressure_start_time")),
                "surgical_treatment": (payload.get("surgical_treatment") or "").strip() or None,
                "surgical_treatment_start_time": _optional_datetime(payload.get("surgical_treatment_start_time")),
                "mild_hypothermia": (payload.get("mild_hypothermia") or "").strip() or None,
                "mild_hypothermia_start_time": _optional_datetime(payload.get("mild_hypothermia_start_time")),
                "brain_protection_drugs": (payload.get("brain_protection_drugs") or "").strip() or None,
                "brain_protection_start_time": _optional_datetime(payload.get("brain_protection_start_time")),
                "antiepileptic_drugs": (payload.get("antiepileptic_drugs") or "").strip() or None,
                "antiepileptic_start_time": _optional_datetime(payload.get("antiepileptic_start_time")),
                "surgery_methods": (payload.get("surgery_methods") or "").strip() or None,
                "surgery_start_time": _optional_datetime(payload.get("surgery_start_time")),
                "chest_fixation": (payload.get("chest_fixation") or "").strip() or None,
                "chest_fixation_start_time": _optional_datetime(payload.get("chest_fixation_start_time")),
                "airway_control": (payload.get("airway_control") or "").strip() or None,
                "airway_control_start_time": _optional_datetime(payload.get("airway_control_start_time")),
                "oxygen_support": (payload.get("oxygen_support") or "").strip() or None,
                "oxygen_support_start_time": _optional_datetime(payload.get("oxygen_support_start_time")),
                "blood_transfusion": (payload.get("blood_transfusion") or "").strip() or None,
                "blood_transfusion_start_time": _optional_datetime(payload.get("blood_transfusion_start_time")),
                "blood_transfusion_total": (payload.get("blood_transfusion_total") or "").strip() or None,
                "temperature_management": (payload.get("temperature_management") or "").strip() or None,
                "temperature_management_start_time": _optional_datetime(payload.get("temperature_management_start_time")),
            }
            columns = [column for column in values.keys() if column in treatment_columns]
            placeholders = ",".join(["%s"] * len(columns))

            cur.execute(
                f"INSERT INTO treatments ({','.join(columns)}) VALUES ({placeholders})",
                [values[column] for column in columns],
            )
            completed_now = _mark_case_complete_if_ready(cur, patient_id)
        conn.commit()
        return ok({"completed_now": completed_now}, "治疗记录已添加" + ("，已自动转为完整记录" if completed_now else ""))
    finally:
        conn.close()


@records_bp.post("/api/patients/<int:patient_id>/followups")
@require_user
def add_followup(patient_id):
    payload = request.get_json(silent=True) or request.form.to_dict()
    follow_time = (payload.get("follow_time") or "").strip()
    diagnosis_record_id = payload.get("diagnosis_record_id")
    if not follow_time or not diagnosis_record_id:
        return fail("请选择诊断记录并填写随访时间")

    conn = db()
    try:
        with conn.cursor() as cur:
            current_user = _current_user_scope(cur)
            if not current_user:
                return fail("用户不存在", 404)
            if not _patient_is_accessible(cur, patient_id, current_user):
                return fail("病人不存在", 404)

            if _table_exists(cur, "diagnosis_records"):
                cur.execute(
                    "SELECT diagnosis_disease FROM diagnosis_records WHERE id=%s AND patient_id=%s LIMIT 1",
                    (diagnosis_record_id, patient_id),
                )
                diagnosis = cur.fetchone()
                if not diagnosis:
                    return fail("请选择有效的诊断记录")

            cur.execute("SHOW COLUMNS FROM followups")
            followup_columns = {row["Field"] for row in cur.fetchall()}
            values = {
                "patient_id": patient_id,
                "user_id": session.get("user_id"),
                "diagnosis_record_id": diagnosis_record_id,
                "follow_time": follow_time,
                "follow_result": "结构化随访记录",
                "prognosis": _optional_number(payload.get("prognosis")),
                "death_days": _optional_number(payload.get("death_days")),
                "barthel_28d": _optional_number(payload.get("barthel_28d")),
                "ventilator_days": _optional_number(payload.get("ventilator_days")),
                "tracheotomy": _optional_number(payload.get("tracheotomy")),
                "blood_purification": _optional_number(payload.get("blood_purification")),
                "total_cost": _optional_number(payload.get("total_cost")),
                "mods": _optional_number(payload.get("mods")),
                "sepsis": _optional_number(payload.get("sepsis")),
                "pulmonary_infection": _optional_number(payload.get("pulmonary_infection")),
                "icu_days": _optional_number(payload.get("icu_days")),
            }
            columns = [column for column in values.keys() if column in followup_columns]
            placeholders = ",".join(["%s"] * len(columns))

            cur.execute(
                f"INSERT INTO followups ({','.join(columns)}) VALUES ({placeholders})",
                [values[column] for column in columns],
            )
            completed_now = _mark_case_complete_if_ready(cur, patient_id)
        conn.commit()
        return ok({"completed_now": completed_now}, "随访记录已添加" + ("，已自动转为完整记录" if completed_now else ""))
    finally:
        conn.close()


@records_bp.post("/api/patients/<int:patient_id>/assessments")
@require_user
def add_assessment(patient_id):
    payload = request.get_json(silent=True) or request.form.to_dict()
    assessment_time = (payload.get("assessment_time") or "").strip()
    diagnosis_record_id = payload.get("diagnosis_record_id")
    if not assessment_time or not diagnosis_record_id:
        return fail("请选择诊断记录并填写评估时间")

    conn = db()
    try:
        with conn.cursor() as cur:
            current_user = _current_user_scope(cur)
            if not current_user:
                return fail("用户不存在", 404)
            if not _patient_is_accessible(cur, patient_id, current_user):
                return fail("病人不存在", 404)
            if not _table_exists(cur, "assessments"):
                return fail("评估记录表不存在，请先执行 ensure_assessment_fields.sql")

            cur.execute(
                "SELECT diagnosis_disease FROM diagnosis_records WHERE id=%s AND patient_id=%s LIMIT 1",
                (diagnosis_record_id, patient_id),
            )
            if not cur.fetchone():
                return fail("请选择有效的诊断记录")

            cur.execute("SHOW COLUMNS FROM assessments")
            assessment_columns = {row["Field"] for row in cur.fetchall()}
            values = {
                "patient_id": patient_id,
                "user_id": session.get("user_id"),
                "diagnosis_record_id": diagnosis_record_id,
                "assessment_time": assessment_time,
                "temperature": _optional_decimal(payload.get("temperature")),
                "respiration": _optional_decimal(payload.get("respiration")),
                "systolic_bp": _optional_decimal(payload.get("systolic_bp")),
                "diastolic_bp": _optional_decimal(payload.get("diastolic_bp")),
                "heart_rate": _optional_decimal(payload.get("heart_rate")),
                "shock_index": _optional_decimal(payload.get("shock_index")),
                "oxygen_partial_pressure": _optional_decimal(payload.get("oxygen_partial_pressure")),
                "oxygen_concentration": _optional_decimal(payload.get("oxygen_concentration")),
                "sofa_score": _optional_decimal(payload.get("sofa_score")),
                "apache_ii_score": _optional_decimal(payload.get("apache_ii_score")),
                "barthel_score": _optional_decimal(payload.get("barthel_score")),
                "mods_score": _optional_decimal(payload.get("mods_score")),
                "gcs_score": _optional_decimal(payload.get("gcs_score")),
                "nihss_score": _optional_decimal(payload.get("nihss_score")),
                "cerebral_hernia": _optional_number(payload.get("cerebral_hernia")),
                "oxygen_saturation": _optional_decimal(payload.get("oxygen_saturation")),
                "ais_score": _optional_decimal(payload.get("ais_score")),
                "pain_score": _optional_decimal(payload.get("pain_score")),
                "iss_score": _optional_decimal(payload.get("iss_score")),
                "created_at": datetime.now(),
            }
            columns = [column for column in values.keys() if column in assessment_columns]
            placeholders = ",".join(["%s"] * len(columns))

            cur.execute(
                f"INSERT INTO assessments ({','.join(columns)}) VALUES ({placeholders})",
                [values[column] for column in columns],
            )
            completed_now = _mark_case_complete_if_ready(cur, patient_id)
        conn.commit()
        return ok({"completed_now": completed_now}, "评估记录已添加" + ("，已自动转为完整记录" if completed_now else ""))
    finally:
        conn.close()
